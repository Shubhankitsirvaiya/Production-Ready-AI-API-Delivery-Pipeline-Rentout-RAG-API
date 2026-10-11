"""Generate the RentOut demo knowledge-base documents.

The chatbot needs realistic source material before ingestion can demonstrate RAG
properly. This script creates detailed `.docx` policy documents with consistent
metadata fields at the top of each file. The ingestion pipeline later reads those
fields and attaches them to every vector chunk for filtering and citations.

Run:
    python scripts/generate_rentout_docs.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "app" / "data" / "source_docs"
OUT_DIR.mkdir(parents=True, exist_ok=True)


DOCUMENTS = [
    {
        "file": "rentout_rental_agreement.docx",
        "title": "RentOut Rental Agreement",
        "document_type": "agreement",
        "department": "legal",
        "audience": "customer",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This agreement defines the baseline obligations between RentOut and a customer renting residential furniture, decor, office furniture, or bundled room packages.",
                "It applies to standard monthly rentals, flexible subscription rentals, short-term staging rentals, and approved corporate rentals unless a signed master services agreement overrides these terms.",
            ],
            "Customer Responsibilities": [
                "Customers must provide accurate delivery information, maintain safe access to the delivery location, use furniture for ordinary indoor purposes, and notify RentOut within forty-eight hours of delivery if an item arrives damaged or incorrect.",
                "Customers must not sell, sublease, pledge, modify, repaint, reupholster, or permanently alter RentOut furniture. Temporary protective covers are allowed when they do not damage fabric, wood, leather, metal, or electrical components.",
                "Customers are responsible for keeping items in climate-controlled indoor spaces unless the product page explicitly identifies the item as outdoor-rated.",
            ],
            "Rental Term and Renewal": [
                "The initial rental term begins on the confirmed delivery date and continues for the selected plan duration. Monthly plans renew automatically until cancelled through the approved cancellation process.",
                "Fixed-term plans may convert to month-to-month at the end of the minimum term when the account is current and the rented items remain available for continued rental.",
            ],
            "Ownership and Restrictions": [
                "All rented items remain property of RentOut unless the customer exercises an approved purchase option. The customer receives temporary possession, not ownership.",
                "RentOut may request confirmation of item location when fraud, loss, unpaid balance, or unauthorized relocation is suspected.",
            ],
            "Examples": [
                "A customer renting a dining set for six months may use it for normal household meals but may not sand, stain, or resize the table.",
                "A customer moving to a new apartment must request relocation support instead of independently moving large rented items when the plan requires RentOut pickup and redelivery.",
            ],
            "Escalation Rules": [
                "Escalate to Legal Operations when a customer disputes ownership, refuses return access, reports theft, or claims a third party has taken possession of RentOut property.",
                "Escalate to Support Leadership when a customer asks for exception approval beyond standard cancellation, relocation, or damage rules.",
            ],
        },
    },
    {
        "file": "deposit_and_refund_policy.docx",
        "title": "Deposit and Refund Policy",
        "document_type": "policy",
        "department": "finance",
        "audience": "customer",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy explains when RentOut collects deposits, how deposits are applied, and when refunds are issued after pickup and inspection.",
                "The policy is designed to make deposit outcomes predictable while protecting RentOut from unpaid balances, missing items, and chargeable damage.",
            ],
            "Deposit Collection": [
                "A refundable deposit may be collected before delivery based on customer profile, rental value, plan duration, location, and prior account history.",
                "Corporate customers may satisfy deposit requirements through an approved corporate guarantee or purchase order if Finance has approved the account terms.",
            ],
            "Refund Eligibility": [
                "Deposits are refundable after all rented items are returned, inspected, and reconciled against the account ledger.",
                "Refunds may be reduced by unpaid rent, late fees, missed pickup fees, missing accessories, cleaning charges, repair costs, or replacement costs for items beyond economical repair.",
            ],
            "Inspection Timeline": [
                "RentOut targets deposit review within seven business days after successful pickup. Complex claims, missing item investigations, or third-party repair estimates may extend the review period.",
                "Customers should receive a refund decision or claim notice once inspection is complete.",
            ],
            "Customer Disputes": [
                "Customers may dispute a deposit deduction by providing photos, delivery notes, chat records, or other evidence within ten business days of the deduction notice.",
                "Support agents should compare the dispute evidence against delivery photos, pickup photos, warehouse inspection notes, and maintenance records.",
            ],
            "Examples": [
                "If a sofa is returned with normal cushion compression and no stains, no damage deduction should be applied.",
                "If a dining chair is missing two legs due to customer disassembly, repair or replacement costs may be deducted from the deposit.",
            ],
            "Escalation Rules": [
                "Escalate refund disputes above 500 USD or involving legal threats to Finance Operations.",
                "Escalate claims involving suspected fraud, chargebacks, or identity mismatch to Risk Operations.",
            ],
        },
    },
    {
        "file": "delivery_and_pickup_policy.docx",
        "title": "Delivery and Pickup Policy",
        "document_type": "policy",
        "department": "operations",
        "audience": "customer",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy defines how RentOut schedules, completes, reschedules, and documents delivery and pickup appointments.",
                "The goal is to protect customer experience while giving operations teams clear rules for access, failed appointments, and safety limits.",
            ],
            "Scheduling Rules": [
                "Delivery and pickup windows are confirmed through email, SMS, or customer portal notification. Customers must provide a valid contact number and accurate access details.",
                "RentOut may require elevator reservations, building certificates, loading dock instructions, gate codes, parking instructions, or proof of permission for managed buildings.",
            ],
            "Customer Readiness": [
                "Customers must ensure the delivery path is clear, pets are secured, and the room is ready for placement. RentOut teams are not responsible for moving non-RentOut furniture unless the service is explicitly purchased.",
                "Delivery teams may refuse entry or reschedule if the location is unsafe, access is blocked, or the requested placement could damage property or create a safety risk.",
            ],
            "Missed Appointments": [
                "A missed appointment occurs when the customer, authorized representative, building access, or required documentation is unavailable during the confirmed window.",
                "Missed appointment fees may apply when RentOut cannot complete delivery or pickup due to customer-controlled factors.",
            ],
            "Pickup Completion": [
                "Pickup is complete only when all scheduled items and accessories are collected and documented. Partial pickups remain open until missing items are resolved.",
                "Drivers should photograph item condition and note visible damage, missing accessories, access issues, and customer comments.",
            ],
            "Examples": [
                "If a customer forgets to reserve the service elevator and the building denies access, the appointment may be treated as missed.",
                "If severe weather prevents safe delivery, RentOut should reschedule without applying a customer missed appointment fee.",
            ],
            "Escalation Rules": [
                "Escalate to Dispatch when access constraints require special equipment, additional labor, or a building-specific delivery plan.",
                "Escalate to Safety Operations when stairs, narrow access, hazardous materials, aggressive animals, or blocked exits create risk.",
            ],
        },
    },
    {
        "file": "damage_and_repair_policy.docx",
        "title": "Furniture Damage and Repair Policy",
        "document_type": "policy",
        "department": "warehouse",
        "audience": "support_agent",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy defines how RentOut distinguishes ordinary wear from chargeable damage and how repair decisions are documented.",
                "The policy supports consistent customer treatment and protects inventory quality across rental cycles.",
            ],
            "Ordinary Wear": [
                "Ordinary wear includes light cushion compression, minor fabric softening, small surface marks consistent with normal use, and hardware loosening that can be corrected during routine maintenance.",
                "Ordinary wear should not create customer charges when the item was used as intended and returned with all required accessories.",
            ],
            "Chargeable Damage": [
                "Chargeable damage includes deep stains, pet urine, burns, broken frames, missing parts, water swelling, unauthorized modifications, pest contamination, cracked glass, and structural damage caused by misuse.",
                "Damage charges must be supported by photos, inspection notes, repair estimates, or replacement records.",
            ],
            "Repair Decisioning": [
                "Warehouse teams classify each item as return-to-stock, clean-only, repairable, parts-required, quarantine, or retire-and-replace.",
                "Items with suspected pest contamination, mold, biohazard exposure, or unsafe electrical components must be quarantined before any further handling.",
            ],
            "Customer Evidence Review": [
                "Support agents should compare pickup photos, delivery photos, warehouse notes, and customer-provided evidence before confirming a charge.",
                "If the customer reported damage within forty-eight hours of delivery and evidence supports pre-existing damage, the claim should not be billed to the customer.",
            ],
            "Examples": [
                "A dining table with faint surface scuffs may be ordinary wear. A dining table with a deep heat burn from a pan may be chargeable damage.",
                "A sofa with removable cushion wrinkles may be ordinary wear. A sofa with pet urine odor and staining may require cleaning or replacement charges.",
            ],
            "Escalation Rules": [
                "Escalate any claim over 750 USD, pest contamination, biohazard handling, or disputed pre-existing damage to Warehouse Quality leadership.",
                "Escalate potential insurance claims to Risk Operations.",
            ],
        },
    },
    {
        "file": "subscription_billing_policy.docx",
        "title": "Subscription Plans and Billing Policy",
        "document_type": "policy",
        "department": "finance",
        "audience": "customer",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy explains how RentOut bills monthly furniture subscriptions, fixed-term plans, plan upgrades, plan downgrades, and failed payments.",
                "It helps agents answer questions consistently while protecting revenue recognition and customer transparency.",
            ],
            "Billing Cycle": [
                "Monthly subscription charges are billed in advance for the upcoming rental period. The billing date is based on the delivery date unless otherwise specified in the customer agreement.",
                "Fixed-term plans may have minimum commitment periods, promotional rates, or early termination terms shown at checkout or in the signed quote.",
            ],
            "Plan Changes": [
                "Customers may request upgrades, item swaps, or plan extensions through the portal or support team. Pricing changes begin when the new item or plan takes effect.",
                "Downgrades may be limited during minimum commitment periods or when the requested change requires pickup, redelivery, or warehouse handling.",
            ],
            "Failed Payments": [
                "When a payment fails, RentOut may retry payment, notify the customer, pause optional services, assess late fees, or begin collections workflows according to account status.",
                "Support agents should never promise fee waivers without checking eligibility and approval rules.",
            ],
            "Corporate Billing": [
                "Corporate rentals may be billed through invoice, purchase order, or central account billing. Corporate terms are governed by the approved account profile.",
                "Support agents should route corporate billing disputes to the assigned account manager when a master services agreement applies.",
            ],
            "Examples": [
                "If a customer swaps a basic desk for an executive desk mid-cycle, the new rate may be prorated from the effective date.",
                "If a payment card expires, the customer remains responsible for unpaid rent even if reminders were not opened.",
            ],
            "Escalation Rules": [
                "Escalate failed payment disputes above 300 USD to Finance Operations.",
                "Escalate corporate invoice mismatches to Account Management and Finance.",
            ],
        },
    },
    {
        "file": "cancellation_early_return_policy.docx",
        "title": "Cancellation and Early Return Policy",
        "document_type": "policy",
        "department": "support",
        "audience": "customer",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy explains how customers cancel future deliveries, end active rentals, and request early returns.",
                "It distinguishes pre-delivery cancellation from active-rental return because the operational cost and customer obligations differ.",
            ],
            "Pre-Delivery Cancellation": [
                "Customers may request cancellation before delivery. Fees may depend on how close the request is to the scheduled route date and whether custom procurement or special handling has already occurred.",
                "If a route has already been loaded or dispatched, cancellation may be treated as a late cancellation.",
            ],
            "Active Rental Return": [
                "Customers may request pickup to end an active rental. Monthly charges continue until the pickup is completed or until another policy-defined end date applies.",
                "Early returns during a minimum commitment period may trigger early termination charges or loss of promotional pricing.",
            ],
            "Refund Handling": [
                "Refunds depend on billing cycle, plan terms, item condition, pickup completion, and open account balances.",
                "Support agents should avoid promising refunds before pickup and inspection are complete unless the policy explicitly allows a pre-approved exception.",
            ],
            "Exceptions": [
                "Exceptions may be considered for documented duplicate orders, RentOut-caused delivery failure, product unavailability, verified medical emergency, or corporate account direction from an authorized approver.",
                "Exceptions must be documented with the reason, approver, and customer communication record.",
            ],
            "Examples": [
                "A customer cancelling three days before a standard delivery may be eligible for no-fee cancellation if the route has not been loaded.",
                "A customer returning a bedroom set two months into a twelve-month promotional term may owe early termination charges.",
            ],
            "Escalation Rules": [
                "Escalate cancellation exceptions outside policy to Support Leadership.",
                "Escalate potential refund commitments above 250 USD to Finance Operations before communicating approval.",
            ],
        },
    },
    {
        "file": "product_care_guide.docx",
        "title": "Product Care Guide",
        "document_type": "guide",
        "department": "support",
        "audience": "customer",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This guide gives customers practical care instructions for rented furniture so items remain safe, clean, and returnable.",
                "Following this guide reduces avoidable damage claims and improves customer experience.",
            ],
            "Fabric Upholstery": [
                "Vacuum fabric upholstery regularly using a soft brush attachment. Blot spills immediately with a clean dry cloth and avoid rubbing, which can spread stains.",
                "Do not use bleach, harsh solvents, steam cleaners, or unapproved stain removers on RentOut fabric items.",
            ],
            "Leather and Faux Leather": [
                "Wipe leather and faux leather with a soft dry or slightly damp cloth. Keep items away from direct heat and prolonged sunlight.",
                "Do not apply oils, conditioners, dyes, or adhesive patches unless RentOut support provides written approval.",
            ],
            "Wood and Veneer": [
                "Use coasters, placemats, and protective pads. Avoid placing hot cookware, wet planters, candles, or adhesive hooks directly on wood surfaces.",
                "Clean with a dry or slightly damp cloth. Excess water can cause swelling, discoloration, or veneer lifting.",
            ],
            "Mattresses and Soft Goods": [
                "Customers should use mattress protectors when provided or recommended. Stains, pest exposure, smoke odor, or biological contamination may make soft goods non-returnable.",
                "Do not remove law labels, covers, zippers, or tags from soft goods.",
            ],
            "Examples": [
                "A water ring from a cup may be preventable by using coasters and may be chargeable if it requires refinishing.",
                "A small removable lint mark on fabric is usually ordinary cleaning, not damage.",
            ],
            "Escalation Rules": [
                "Escalate care questions involving chemical exposure, pest contamination, smoke odor, or suspected mold to Support Leadership.",
                "Escalate safety issues involving electrical items to Maintenance Operations.",
            ],
        },
    },
    {
        "file": "corporate_rental_policy.docx",
        "title": "Corporate Furniture Rental Policy",
        "document_type": "policy",
        "department": "account_management",
        "audience": "corporate",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy governs corporate rentals for offices, relocation programs, model apartments, temporary workspaces, and employee housing.",
                "Corporate rentals may include account-level pricing, purchase orders, named approvers, service-level expectations, and consolidated billing.",
            ],
            "Authorized Requesters": [
                "Only authorized corporate contacts may place orders, approve swaps, request pickups, or modify billing instructions.",
                "Support agents must verify requester authority before making account-level changes.",
            ],
            "Service Scope": [
                "Corporate service may include space planning, bundled room packages, staged delivery waves, installation coordination, and scheduled refresh cycles.",
                "Services outside the approved scope require written change approval from the corporate contact or account manager.",
            ],
            "Billing and Deposits": [
                "Corporate accounts may use invoice billing, purchase order billing, card-on-file billing, or master account billing based on approved terms.",
                "Deposits may be waived or replaced with corporate guarantee terms when approved by Finance.",
            ],
            "Returns and Asset Tracking": [
                "Corporate customers must maintain location records for rented assets and notify RentOut before moving items between sites.",
                "Bulk pickup requires item reconciliation against the corporate asset list.",
            ],
            "Examples": [
                "A relocation provider may rent bedroom and living room sets for employees on six-month assignments.",
                "A startup may rent modular office furniture under a twelve-month corporate plan with quarterly refresh options.",
            ],
            "Escalation Rules": [
                "Escalate unauthorized requester changes to Account Management.",
                "Escalate large corporate disputes, master agreement conflicts, or multi-site pickup failures to the assigned account manager.",
            ],
        },
    },
    {
        "file": "support_escalation_sop.docx",
        "title": "Customer Support Escalation SOP",
        "document_type": "sop",
        "department": "support",
        "audience": "support_agent",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This SOP tells RentOut support agents when to resolve an issue directly and when to escalate to another team.",
                "It is intended to reduce inconsistent handling, missed risk signals, and unnecessary delays.",
            ],
            "Agent-Owned Issues": [
                "Agents may resolve standard plan questions, delivery status questions, basic pickup scheduling, product care guidance, and policy explanations using approved knowledge base content.",
                "Agents should document the customer question, policy basis, action taken, and any promised follow-up.",
            ],
            "Required Escalations": [
                "Escalate legal threats, chargebacks, suspected fraud, safety incidents, pest contamination, refund disputes above threshold, corporate contract conflicts, and media or executive complaints.",
                "Escalate any issue where the customer requests a commitment that exceeds published policy or the agent's authority.",
            ],
            "Escalation Package": [
                "The escalation note should include customer ID, order ID, item IDs, relevant photos, timeline, customer request, policy references, financial exposure, and recommended next action.",
                "If evidence is incomplete, state exactly what is missing rather than delaying escalation indefinitely.",
            ],
            "Response Time Expectations": [
                "Urgent safety and access issues should be routed immediately. Standard refund or exception reviews should be routed by the end of the same business day.",
                "Agents should avoid repeated handoffs by selecting the correct owning team on the first escalation.",
            ],
            "Examples": [
                "A customer asking how to clean a fabric sofa is agent-owned.",
                "A customer threatening legal action over a 900 USD damage claim requires escalation.",
            ],
            "Escalation Rules": [
                "Escalate to Risk for fraud, chargeback, identity mismatch, theft, or insurance matters.",
                "Escalate to Warehouse Quality for disputed damage classification or contamination.",
                "Escalate to Finance for refunds, billing disputes, and deposit deductions above threshold.",
            ],
        },
    },
    {
        "file": "privacy_customer_data_policy.docx",
        "title": "Privacy and Customer Data Handling Policy",
        "document_type": "policy",
        "department": "security",
        "audience": "internal",
        "version": "v1.0",
        "effective_date": "2026-01-01",
        "sections": {
            "Purpose": [
                "This policy defines how RentOut teams handle customer personal data in support, operations, billing, delivery, and account management workflows.",
                "It protects customer trust and limits unnecessary exposure of sensitive information.",
            ],
            "Permitted Data Use": [
                "Customer data may be used to process orders, schedule delivery, verify account ownership, manage billing, resolve disputes, prevent fraud, and comply with legal obligations.",
                "Employees should access only the information needed for the current task.",
            ],
            "Restricted Data": [
                "Payment card numbers, identity documents, access codes, building entry instructions, and sensitive complaint details must be handled with extra care.",
                "Support agents must not paste full payment details, identity documents, or access codes into unsecured notes or external tools.",
            ],
            "Customer Verification": [
                "Before discussing account details, agents should verify the requester using approved identifiers such as account email, order ID, phone verification, or corporate authorized contact status.",
                "Corporate account changes require requester authority verification.",
            ],
            "Retention and Logs": [
                "Operational notes should be factual, minimal, and relevant. Do not include unnecessary personal commentary or unrelated sensitive details.",
                "Audit logs should record the action taken, policy basis, and responsible user without duplicating sensitive content unless required.",
            ],
            "Examples": [
                "It is acceptable to note that a customer requested pickup due to relocation. It is not acceptable to include unrelated medical details unless required for an approved exception.",
                "A building gate code should be stored only in approved delivery instructions, not copied into free-form chat notes.",
            ],
            "Escalation Rules": [
                "Escalate suspected data exposure, unauthorized access, or account takeover concerns to Security Operations.",
                "Escalate legal data requests to Legal Operations.",
            ],
        },
    },
]


def main() -> None:
    """Create every DOCX file declared in DOCUMENTS."""
    for spec in DOCUMENTS:
        build_document(spec)
    print(f"Generated {len(DOCUMENTS)} DOCX files in {OUT_DIR}")


def build_document(spec: dict) -> None:
    """Build one Word document from a dictionary specification.

    spec contains both metadata fields and section text. Keeping the source data
    structured makes it easy to add more policies without editing formatting
    logic.
    """
    doc = Document()
    set_styles(doc)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(spec["title"])
    run.font.color.rgb = RGBColor(0, 109, 119)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("RentOut Furniture Rental Operations").italic = True

    # These header fields are intentionally machine-readable. app/scripts_api.py looks for `Key: Value` lines near the top of the document and turns them into retrieval metadata.
    doc.add_paragraph(f"Title: {spec['title']}")
    doc.add_paragraph(f"Document Type: {spec['document_type']}")
    doc.add_paragraph(f"Department: {spec['department']}")
    doc.add_paragraph(f"Audience: {spec['audience']}")
    doc.add_paragraph(f"Version: {spec['version']}")
    doc.add_paragraph(f"Effective Date: {spec['effective_date']}")

    add_callout(
        doc,
        "Document Use",
        "This document is part of the RentOut RAG chatbot corpus. It contains detailed operational language, metadata, examples, and escalation rules so retrieval can use both semantic meaning and structured source attributes.",
    )

    # Body sections provide natural-language content for semantic retrieval.
    for section, paragraphs in spec["sections"].items():
        doc.add_heading(section, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)

    # The metadata table is human-readable. The header lines above are what the ingestion code parses automatically.
    doc.add_heading("Metadata for Retrieval", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Value"
    for field in ["title", "document_type", "department", "audience", "version", "effective_date"]:
        row = table.add_row().cells
        row[0].text = field
        row[1].text = str(spec[field])

    doc.add_heading("Version History", level=1)
    doc.add_paragraph(f"{spec['version']} - Initial RentOut RAG deployment corpus version, effective {spec['effective_date']}.")

    doc.save(OUT_DIR / spec["file"])


def set_styles(doc: Document) -> None:
    """Apply simple corporate styling so the generated docs look realistic."""
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.color.rgb = RGBColor(0, 109, 119)


def add_callout(doc: Document, label: str, body: str) -> None:
    """Add a bold labeled paragraph used for document context notes."""
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor(0, 109, 119)
    paragraph.add_run(body)


if __name__ == "__main__":
    main()

